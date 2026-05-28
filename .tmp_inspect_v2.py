from openpyxl import load_workbook
import re
p=r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT/analysis/mod_new_pm_cn_grouped_v2.xlsx"
ws=load_workbook(p).active
u=[]
e=[]
for i,row in enumerate(ws.iter_rows(min_row=2, values_only=True),start=2):
    txt=' | '.join([str(row[j] or '') for j in [0,1,2,3,5]])
    if '未匹配' in txt:
        u.append((i,row[0],row[1],row[2]))
    if re.search(r'[A-Za-z_]{3,}', txt):
        e.append((i,row[0],row[1],row[2],row[3],row[5]))
print('unmatched',len(u))
for x in u: print(x)
print('english',len(e))
for x in e[:25]: print(x[:4])
