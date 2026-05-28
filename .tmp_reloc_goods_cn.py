import re
from pathlib import Path
from openpyxl import load_workbook

mod_root=Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
van_loc=Path(r"D:/Steam/steamapps/common/Victoria 3/game/localization/simp_chinese")
mod_loc=mod_root/'localization'/'simp_chinese'


def parse_loc(folder):
    d={}
    if not folder.exists(): return d
    for f in folder.rglob('*.yml'):
        t=f.read_text(encoding='utf-8-sig',errors='ignore')
        for line in t.splitlines():
            s=line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'): continue
            m=re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$',s)
            if m:
                d[m.group(1)]=m.group(2).replace('\\"','"').strip()
    return d

loc={}
loc.update(parse_loc(van_loc))
loc.update(parse_loc(mod_loc))

inp=mod_root/'analysis'/'mod_new_pm_cn_grouped_v4.xlsx'
out=mod_root/'analysis'/'mod_new_pm_cn_grouped_v5_原版商品中文.xlsx'
wb=load_workbook(inp)
ws=wb.active

for r in ws.iter_rows(min_row=2):
    for idx in [3,5]:
        v=r[idx].value
        if not isinstance(v,str) or not v: continue
        parts=v.split('；')
        new=[]
        for p in parts:
            if '*' not in p:
                new.append(p); continue
            name,qty=p.rsplit('*',1)
            key=name
            # try reverse map from existing chinese to key is hard; detect alias patterns with fallback list
            # if already pure chinese keep; if looks like key use localization
            if re.match(r'^[A-Za-z0-9_]+$', name):
                key=name
                name=loc.get(key,key)
            new.append(f"{name}*{qty}")
        r[idx].value='；'.join(new)

wb.save(out)
print(out)
