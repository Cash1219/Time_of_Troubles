from openpyxl import load_workbook
from pathlib import Path
p=Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT/analysis/mod_new_pm_cn_grouped_v3.xlsx")
wb=load_workbook(p)
ws=wb.active
for r in ws.iter_rows(min_row=2):
    # 建筑
    if r[0].value == 'building_railway':
        r[0].value = '铁路'
    # 生产方式
    if r[2].value in ('pm_no_magic','������ħ','不启用'):
        r[2].value = '不启用'
    if r[2].value in ('pm_railway_luotuomu','�����','洛托姆货运'):
        r[2].value = '洛托姆货运'
    if r[2].value in ('pm_railway_transport','����ķ����','反重力附魔'):
        r[2].value = '反重力附魔'
out=Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT/analysis/mod_new_pm_cn_grouped_v4.xlsx")
wb.save(out)
print(out)
