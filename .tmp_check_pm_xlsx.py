from openpyxl import load_workbook
from pathlib import Path
import re
p=Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT/analysis/mod_new_production_methods_value_中文精简_按建筑PMGPM.xlsx")
wb=load_workbook(p)
ws=wb.active
unmatched=[]
english=[]
for i,row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    b,pmg,pm,ins,_,outs,_,_,_=row
    txt=' | '.join([str(x or '') for x in [b,pmg,pm,ins,outs]])
    if '未匹配' in txt:
        unmatched.append((i,b,pmg,pm))
    if re.search(r'[A-Za-z_]{3,}', txt):
        english.append((i,b,pmg,pm,ins,outs))
print('unmatched',len(unmatched))
print('english',len(english))
for r in unmatched[:20]:
    print('U',r)
for r in english[:20]:
    print('E',r[:4])
