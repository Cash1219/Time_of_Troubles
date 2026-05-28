import re, json
from pathlib import Path
from openpyxl import Workbook

root = Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
mod_pm_dir = root / "common" / "production_methods"
van_pm_dir = Path(r"D:/Steam/steamapps/common/Victoria 3/game/common/production_methods")
van_goods = Path(r"D:/Steam/steamapps/common/Victoria 3/game/common/goods/00_goods.txt")
mod_goods_dir = root / "common" / "goods"
loc_dir = root / "localization" / "simp_chinese"


def parse_top_blocks(text):
    blocks = {}
    i, n = 0, len(text)
    pat = re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{')
    while True:
        m = pat.search(text, i)
        if not m:
            break
        name = m.group(1)
        brace_start = m.end() - 1
        depth = 0
        j = brace_start
        while j < n:
            c = text[j]
            if c == '{':
                depth += 1
            elif c == '}':
                depth -= 1
                if depth == 0:
                    blocks[name] = text[brace_start + 1:j]
                    i = j + 1
                    break
            j += 1
        else:
            break
    return blocks


def parse_loc_files(folder: Path):
    loc = {}
    for f in folder.rglob('*.yml'):
        txt = f.read_text(encoding='utf-8-sig', errors='ignore')
        for line in txt.splitlines():
            s = line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'):
                continue
            m = re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$', s)
            if m:
                key, val = m.group(1), m.group(2)
                val = val.replace('\\"', '"').strip()
                loc[key] = val
    return loc


def load_goods_costs():
    costs = {}
    for k, body in parse_top_blocks(van_goods.read_text(encoding='utf-8-sig', errors='ignore')).items():
        m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', body)
        if m:
            costs[k] = float(m.group(1))
    if mod_goods_dir.exists():
        for f in mod_goods_dir.glob('*.txt'):
            for k, body in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
                m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', body)
                if m:
                    costs[k] = float(m.group(1))
    return costs


def extract_io(pm_body):
    ins, outs = {}, {}
    for m in re.finditer(r'(?m)^\s*(?!#)goods_input_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', pm_body):
        g, v = m.group(1), float(m.group(2))
        ins[g] = ins.get(g, 0.0) + v
    for m in re.finditer(r'(?m)^\s*(?!#)goods_output_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', pm_body):
        g, v = m.group(1), float(m.group(2))
        outs[g] = outs.get(g, 0.0) + v
    return ins, outs

loc = parse_loc_files(loc_dir)
costs = load_goods_costs()

van_keys = set()
for f in van_pm_dir.glob('*.txt'):
    van_keys.update(parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).keys())

mod_pms = []
for f in mod_pm_dir.glob('*.txt'):
    for k, body in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
        if k not in van_keys:
            mod_pms.append((k, f.name, body))

rows = []
for pm_key, src, body in mod_pms:
    ins, outs = extract_io(body)
    in_total = sum(q * costs.get(g, 0.0) for g, q in ins.items())
    out_total = sum(q * costs.get(g, 0.0) for g, q in outs.items())
    net = out_total - in_total
    roi = (net / in_total) if in_total else None

    pm_cn = loc.get(pm_key, pm_key)

    in_detail = '；'.join([
        f"{loc.get(g,g)}({g})*{q:g}" for g, q in sorted(ins.items())
    ])
    out_detail = '；'.join([
        f"{loc.get(g,g)}({g})*{q:g}" for g, q in sorted(outs.items())
    ])

    rows.append((pm_key, pm_cn, src, in_detail, in_total, out_detail, out_total, net, roi))

rows.sort(key=lambda x: x[0])

out_xlsx = root / 'analysis' / 'mod_new_production_methods_value_中文.xlsx'
out_csv = root / 'analysis' / 'mod_new_production_methods_value_中文.csv'

wb = Workbook()
ws = wb.active
ws.title = '新增生产方式收益(中文)'
headers = ['生产方式key','生产方式中文名','来源文件','商品投入类型与数量','商品投入总价格','商品产出类型与数量','商品产出总价格','商品收益净值','收益率']
ws.append(headers)
for r in rows:
    ws.append(list(r))

for c in ws[1]:
    c.font = c.font.copy(bold=True)

for row in ws.iter_rows(min_row=2, min_col=5, max_col=8):
    for c in row:
        c.number_format = '0.00'
for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
    for c in row:
        c.number_format = '0.00%'

for i,w in {1:34,2:24,3:24,4:68,5:14,6:68,7:14,8:14,9:10}.items():
    ws.column_dimensions[chr(64+i)].width = w

wb.save(out_xlsx)

import csv
with out_csv.open('w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(headers)
    for r in rows:
        key, cn, src, ind, it, oud, ot, net, roi = r
        w.writerow([key,cn,src,ind,f"{it:.2f}",oud,f"{ot:.2f}",f"{net:.2f}",'' if roi is None else f"{roi:.4f}"])

missing_pm_cn = [k for k,cn,_,_,_,_,_,_,_ in rows if cn == k]
print(json.dumps({
    'rows': len(rows),
    'xlsx': str(out_xlsx),
    'csv': str(out_csv),
    'missing_pm_cn_count': len(missing_pm_cn),
    'missing_pm_cn_sample': missing_pm_cn[:20]
}, ensure_ascii=False))
