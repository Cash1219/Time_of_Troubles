from openpyxl import load_workbook
import re
p=r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT/analysis/mod_new_pm_cn_grouped_v2.xlsx"
ws=load_workbook(p).active
bad=[]
for i,row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    for j in [0,1,2,3,5]:
        v=str(row[j] or '')
        if 'pm_' in v or 'pmg_' in v or 'building_' in v or re.search(r'\b(magictools|magic_energy|magic_mineral|magic_steels|omocha)\b', v):
            bad.append((i,j+1,v))
print(len(bad))
for x in bad[:30]:
    print(x)
