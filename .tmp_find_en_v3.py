from openpyxl import load_workbook
import re
p=r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT/analysis/mod_new_pm_cn_grouped_v3.xlsx"
ws=load_workbook(p).active
for i,row in enumerate(ws.iter_rows(min_row=2, values_only=True),start=2):
    txt=' | '.join([str(row[j] or '') for j in [0,1,2,3,5]])
    if re.search(r'\b(pm_|pmg_|building_)',txt):
        print(i,row[0],row[1],row[2])
