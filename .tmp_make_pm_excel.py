import re, json
from pathlib import Path
from openpyxl import Workbook

root = Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
mod_pm_dir = root / "common" / "production_methods"
van_pm_dir = Path(r"D:/Steam/steamapps/common/Victoria 3/game/common/production_methods")
van_goods = Path(r"D:/Steam/steamapps/common/Victoria 3/game/common/goods/00_goods.txt")
mod_goods_dir = root / "common" / "goods"


def parse_top_blocks(text):
    blocks = {}
    i, n = 0, len(text)
    pat = re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{')
    while True:
        m = pat.search(text, i)
        if not m:
            break
        name = m.group(1)
        brace_start = m.end() - 1
        depth = 0
        j = brace_start
        while j < n:
            c = text[j]
            if c == '{':
                depth += 1
            elif c == '}':
                depth -= 1
                if depth == 0:
                    blocks[name] = text[brace_start + 1:j]
                    i = j + 1
                    break
            j += 1
        else:
            break
    return blocks


def load_goods_costs():
    costs = {}
    # vanilla base
    for k, body in parse_top_blocks(van_goods.read_text(encoding='utf-8-sig', errors='ignore')).items():
        m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', body)
        if m:
            costs[k] = float(m.group(1))
    # mod goods override/add
    if mod_goods_dir.exists():
        for f in mod_goods_dir.glob('*.txt'):
            for k, body in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
                m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', body)
                if m:
                    costs[k] = float(m.group(1))
    return costs


def extract_io(pm_body):
    ins, outs = {}, {}
    for m in re.finditer(r'(?m)^\s*(?!#)goods_input_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', pm_body):
        g, v = m.group(1), float(m.group(2))
        ins[g] = ins.get(g, 0.0) + v
    for m in re.finditer(r'(?m)^\s*(?!#)goods_output_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', pm_body):
        g, v = m.group(1), float(m.group(2))
        outs[g] = outs.get(g, 0.0) + v
    return ins, outs

# gather vanilla pm keys
van_keys = set()
for f in van_pm_dir.glob('*.txt'):
    van_keys.update(parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).keys())

# collect mod-added PMs only
mod_pms = []
for f in mod_pm_dir.glob('*.txt'):
    for k, body in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
        if k not in van_keys:
            mod_pms.append((k, f.name, body))

costs = load_goods_costs()
rows = []
for key, src, body in mod_pms:
    ins, outs = extract_io(body)
    in_total = sum(q * costs.get(g, 0.0) for g, q in ins.items())
    out_total = sum(q * costs.get(g, 0.0) for g, q in outs.items())
    net = out_total - in_total
    roi = (net / in_total) if in_total else None
    in_detail = '；'.join([f"{g}*{q:g}" for g, q in sorted(ins.items())])
    out_detail = '；'.join([f"{g}*{q:g}" for g, q in sorted(outs.items())])
    rows.append((key, src, in_detail, in_total, out_detail, out_total, net, roi))

rows.sort(key=lambda x: x[0])

out_xlsx = root / 'analysis' / 'mod_new_production_methods_value.xlsx'
out_csv = root / 'analysis' / 'mod_new_production_methods_value.csv'

wb = Workbook()
ws = wb.active
ws.title = '新增生产方式收益'
headers = ['生产方式key','来源文件','商品投入类型与数量','商品投入总价格','商品产出类型与数量','商品产出总价格','商品收益净值','收益率']
ws.append(headers)
for r in rows:
    ws.append(list(r))

for c in ws[1]:
    c.font = c.font.copy(bold=True)

for row in ws.iter_rows(min_row=2, min_col=4, max_col=7):
    for c in row:
        c.number_format = '0.00'
for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
    for c in row:
        c.number_format = '0.00%'

for i,w in {1:36,2:24,3:60,4:14,5:60,6:14,7:14,8:10}.items():
    ws.column_dimensions[chr(64+i)].width = w

wb.save(out_xlsx)

import csv
with out_csv.open('w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(headers)
    for key,src,ind,it,oud,ot,net,roi in rows:
        w.writerow([key,src,ind,f"{it:.2f}",oud,f"{ot:.2f}",f"{net:.2f}",'' if roi is None else f"{roi:.4f}"])

print(json.dumps({
    'rows': len(rows),
    'with_inputs': sum(1 for r in rows if r[2]),
    'with_outputs': sum(1 for r in rows if r[4]),
    'unknown_goods_count': len({g for r in rows for g in ([x.split('*')[0] for x in r[2].split('；') if x] + [x.split('*')[0] for x in r[4].split('；') if x]) if g not in costs}),
    'xlsx': str(out_xlsx),
    'csv': str(out_csv)
}, ensure_ascii=False))
