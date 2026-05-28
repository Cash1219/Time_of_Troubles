import re
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook

root = Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
van_root = Path(r"D:/Steam/steamapps/common/Victoria 3/game")

mod_pm_dir = root/'common'/'production_methods'
van_pm_dir = van_root/'common'/'production_methods'
mod_pmg_dir = root/'common'/'production_method_groups'
van_pmg_dir = van_root/'common'/'production_method_groups'
mod_bld_dir = root/'common'/'buildings'
van_bld_dir = van_root/'common'/'buildings'
loc_dir = root/'localization'/'simp_chinese'
van_goods = van_root/'common'/'goods'/'00_goods.txt'
mod_goods_dir = root/'common'/'goods'


def parse_top_blocks(text):
    blocks = {}
    i = 0
    pat = re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{')
    n = len(text)
    while True:
        m = pat.search(text, i)
        if not m:
            break
        key = m.group(1)
        s = m.end()-1
        d = 0
        j = s
        while j < n:
            c = text[j]
            if c == '{': d += 1
            elif c == '}':
                d -= 1
                if d == 0:
                    blocks[key] = text[s+1:j]
                    i = j+1
                    break
            j += 1
        else:
            break
    return blocks


def read_blocks_from_dir(d):
    out = {}
    if not d.exists():
        return out
    for f in d.glob('*.txt'):
        txt = f.read_text(encoding='utf-8-sig', errors='ignore')
        for k,v in parse_top_blocks(txt).items():
            out[k] = v
    return out


def parse_loc(folder):
    loc = {}
    for f in folder.rglob('*.yml'):
        t = f.read_text(encoding='utf-8-sig', errors='ignore')
        for line in t.splitlines():
            s = line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'):
                continue
            m = re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$', s)
            if m:
                loc[m.group(1)] = m.group(2).replace('\\"','"').strip()
    return loc


def load_costs():
    costs = {}
    for k,b in parse_top_blocks(van_goods.read_text(encoding='utf-8-sig', errors='ignore')).items():
        m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', b)
        if m: costs[k] = float(m.group(1))
    if mod_goods_dir.exists():
        for f in mod_goods_dir.glob('*.txt'):
            for k,b in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
                m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', b)
                if m: costs[k] = float(m.group(1))
    return costs


def extract_io(body):
    ins, outs = {}, {}
    for m in re.finditer(r'(?m)^\s*(?!#)goods_input_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g, v = m.group(1), float(m.group(2)); ins[g] = ins.get(g,0.0)+v
    for m in re.finditer(r'(?m)^\s*(?!#)goods_output_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g, v = m.group(1), float(m.group(2)); outs[g] = outs.get(g,0.0)+v
    return ins, outs

# merged data (mod override vanilla)
pm_blocks = read_blocks_from_dir(van_pm_dir)
pm_blocks.update(read_blocks_from_dir(mod_pm_dir))
pmg_blocks = read_blocks_from_dir(van_pmg_dir)
pmg_blocks.update(read_blocks_from_dir(mod_pmg_dir))
bld_blocks = read_blocks_from_dir(van_bld_dir)
bld_blocks.update(read_blocks_from_dir(mod_bld_dir))

# identify mod-added PM keys
van_pm_keys = set(read_blocks_from_dir(van_pm_dir).keys())
mod_pm_blocks = read_blocks_from_dir(mod_pm_dir)
new_pm_keys = [k for k in mod_pm_blocks.keys() if k not in van_pm_keys]

# PM -> PMG
pm_to_pmg = defaultdict(set)
for pmg, body in pmg_blocks.items():
    for m in re.finditer(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{', body):
        candidate = m.group(1)
        if candidate in pm_blocks:
            pm_to_pmg[candidate].add(pmg)

# PMG -> Building
pmg_to_bld = defaultdict(set)
for bld, body in bld_blocks.items():
    # production_method_groups = { ... }
    m = re.search(r'(?ms)production_method_groups\s*=\s*\{(.*?)\}', body)
    if not m:
        continue
    inside = m.group(1)
    for g in re.findall(r'(?m)^\s*([A-Za-z0-9_]+)\s*$', inside):
        pmg_to_bld[g].add(bld)

loc = parse_loc(loc_dir)
costs = load_costs()

rows = []
for pm in new_pm_keys:
    body = pm_blocks.get(pm, '')
    ins, outs = extract_io(body)
    in_total = sum(q*costs.get(g,0.0) for g,q in ins.items())
    out_total = sum(q*costs.get(g,0.0) for g,q in outs.items())
    net = out_total - in_total
    roi = (net/in_total) if in_total else None

    pmgs = sorted(pm_to_pmg.get(pm) or ['(未匹配PMG)'])
    for pmg in pmgs:
        blds = sorted(pmg_to_bld.get(pmg) or ['(未匹配建筑)'])
        for bld in blds:
            rows.append({
                'bld': bld, 'bld_cn': loc.get(bld,bld),
                'pmg': pmg, 'pmg_cn': loc.get(pmg,pmg),
                'pm': pm, 'pm_cn': loc.get(pm,pm),
                'in': '；'.join([f"{loc.get(g,g)}*{q:g}" for g,q in sorted(ins.items())]),
                'in_total': in_total,
                'out': '；'.join([f"{loc.get(g,g)}*{q:g}" for g,q in sorted(outs.items())]),
                'out_total': out_total,
                'net': net,
                'roi': roi,
            })

rows.sort(key=lambda r: (r['bld_cn'], r['pmg_cn'], r['pm_cn'], r['pm']))

out = root/'analysis'/'mod_new_production_methods_value_中文精简_按建筑PMGPM.xlsx'
wb = Workbook()
ws = wb.active
ws.title = '建筑-组-生产方式'
ws.append(['建筑','生产方式组','生产方式','商品投入（中文）','商品投入总价格','商品产出（中文）','商品产出总价格','商品收益净值','收益率'])
for r in rows:
    ws.append([r['bld_cn'], r['pmg_cn'], r['pm_cn'], r['in'], r['in_total'], r['out'], r['out_total'], r['net'], r['roi']])
for c in ws[1]:
    c.font = c.font.copy(bold=True)
for row in ws.iter_rows(min_row=2, min_col=5, max_col=8):
    for c in row: c.number_format='0.00'
for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
    for c in row: c.number_format='0.00%'
for i,w in {1:24,2:24,3:24,4:62,5:14,6:62,7:14,8:14,9:10}.items():
    ws.column_dimensions[chr(64+i)].width = w
wb.save(out)
print(out)
print(f"rows={len(rows)}, new_pm={len(new_pm_keys)}")
