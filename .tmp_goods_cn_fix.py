import re
from pathlib import Path
from openpyxl import load_workbook

mod_root=Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
van_loc=Path(r"D:/Steam/steamapps/common/Victoria 3/game/localization/simp_chinese")
mod_loc=mod_root/'localization'/'simp_chinese'
src=mod_root/'analysis'/'mod_new_pm_cn_grouped_v6_顺序修正.xlsx'
out=mod_root/'analysis'/'mod_new_pm_cn_grouped_v7_投入产出全中文.xlsx'

def parse_loc(folder):
    d={}
    if not folder.exists(): return d
    for f in folder.rglob('*.yml'):
        t=f.read_text(encoding='utf-8-sig', errors='ignore')
        for line in t.splitlines():
            s=line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'): continue
            m=re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$', s)
            if m: d[m.group(1)] = m.group(2).replace('\\"','"').strip()
    return d

loc={}
loc.update(parse_loc(van_loc))
loc.update(parse_loc(mod_loc))

alias={
 'magictools': loc.get('magic_tools','魔导器'),
 'magic_mineral': loc.get('magic_mineral','魔法矿石'),
 'magic_energy': loc.get('magic_energy','魔力'),
 'magic_steels': loc.get('magic_steels','超冶金结构'),
 'omocha': loc.get('omocha','周边'),
}

def cn_goods(name):
    if re.match(r'^[A-Za-z0-9_]+$', name):
        return loc.get(name) or alias.get(name) or name
    return name

wb=load_workbook(src)
ws=wb.active
for row in ws.iter_rows(min_row=2):
    for col in [4,6]:  # D,F
        v=row[col-1].value
        if not isinstance(v,str) or not v:
            continue
        parts=v.split('；')
        new=[]
        for p in parts:
            p=p.strip()
            if not p:
                continue
            if '*' in p:
                n,q=p.rsplit('*',1)
                n=cn_goods(n.strip())
                new.append(f"{n}*{q.strip()}")
            else:
                new.append(cn_goods(p))
        row[col-1].value='；'.join(new)

wb.save(out)
print(out)
