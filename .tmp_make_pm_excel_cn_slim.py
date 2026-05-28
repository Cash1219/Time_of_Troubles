import re
from pathlib import Path
from openpyxl import Workbook

root = Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
mod_pm_dir = root / "common" / "production_methods"
van_pm_dir = Path(r"D:/Steam/steamapps/common/Victoria 3/game/common/production_methods")
van_goods = Path(r"D:/Steam/steamapps/common/Victoria 3/game/common/goods/00_goods.txt")
mod_goods_dir = root / "common" / "goods"
loc_dir = root / "localization" / "simp_chinese"

def parse_top_blocks(text):
    blocks = {}
    i, n = 0, len(text)
    pat = re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{')
    while True:
        m = pat.search(text, i)
        if not m:
            break
        name = m.group(1)
        s = m.end()-1
        d = 0
        j = s
        while j < n:
            if text[j] == '{': d += 1
            elif text[j] == '}':
                d -= 1
                if d == 0:
                    blocks[name] = text[s+1:j]
                    i = j+1
                    break
            j += 1
        else:
            break
    return blocks

def parse_loc(folder):
    loc = {}
    for f in folder.rglob('*.yml'):
        t = f.read_text(encoding='utf-8-sig', errors='ignore')
        for line in t.splitlines():
            s = line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'):
                continue
            m = re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$', s)
            if m:
                loc[m.group(1)] = m.group(2).replace('\\"','"').strip()
    return loc

def load_costs():
    costs = {}
    for k,b in parse_top_blocks(van_goods.read_text(encoding='utf-8-sig', errors='ignore')).items():
        m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', b)
        if m: costs[k] = float(m.group(1))
    if mod_goods_dir.exists():
        for f in mod_goods_dir.glob('*.txt'):
            for k,b in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
                m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', b)
                if m: costs[k] = float(m.group(1))
    return costs

def extract_io(body):
    ins, outs = {}, {}
    for m in re.finditer(r'(?m)^\s*(?!#)goods_input_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g, v = m.group(1), float(m.group(2)); ins[g] = ins.get(g,0.0)+v
    for m in re.finditer(r'(?m)^\s*(?!#)goods_output_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g, v = m.group(1), float(m.group(2)); outs[g] = outs.get(g,0.0)+v
    return ins, outs

loc = parse_loc(loc_dir)
costs = load_costs()

van_keys = set()
for f in van_pm_dir.glob('*.txt'):
    van_keys.update(parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).keys())

rows = []
for f in mod_pm_dir.glob('*.txt'):
    for pm_key, body in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
        if pm_key in van_keys:
            continue
        ins, outs = extract_io(body)
        in_total = sum(q*costs.get(g,0.0) for g,q in ins.items())
        out_total = sum(q*costs.get(g,0.0) for g,q in outs.items())
        net = out_total - in_total
        roi = (net / in_total) if in_total else None
        pm_cn = loc.get(pm_key, pm_key)
        in_detail = '；'.join([f"{loc.get(g,g)}*{q:g}" for g,q in sorted(ins.items())])
        out_detail = '；'.join([f"{loc.get(g,g)}*{q:g}" for g,q in sorted(outs.items())])
        rows.append((pm_key, pm_cn, in_detail, in_total, out_detail, out_total, net, roi))

rows.sort(key=lambda x:x[0])

out = root / 'analysis' / 'mod_new_production_methods_value_中文精简.xlsx'
wb = Workbook()
ws = wb.active
ws.title = '生产方式收益'
ws.append(['生产方式key','生产方式中文名','商品投入（中文）','商品投入总价格','商品产出（中文）','商品产出总价格','商品收益净值','收益率'])
for r in rows: ws.append(list(r))
for c in ws[1]: c.font = c.font.copy(bold=True)
for row in ws.iter_rows(min_row=2, min_col=4, max_col=7):
    for c in row: c.number_format = '0.00'
for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
    for c in row: c.number_format = '0.00%'
for i,w in {1:34,2:24,3:62,4:14,5:62,6:14,7:14,8:10}.items():
    ws.column_dimensions[chr(64+i)].width = w
wb.save(out)
print(out)
