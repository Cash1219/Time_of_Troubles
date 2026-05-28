import re
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook

root = Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
van_root = Path(r"D:/Steam/steamapps/common/Victoria 3/game")

loc_dir = root/'localization'/'simp_chinese'
mod_pm_dir = root/'common'/'production_methods'
van_pm_dir = van_root/'common'/'production_methods'
mod_pmg_dir = root/'common'/'production_method_groups'
van_pmg_dir = van_root/'common'/'production_method_groups'
mod_bld_dir = root/'common'/'buildings'
van_bld_dir = van_root/'common'/'buildings'
van_goods = van_root/'common'/'goods'/'00_goods.txt'
mod_goods_dir = root/'common'/'goods'


def iter_files(d):
    if not d.exists():
        return
    for p in d.rglob('*'):
        if p.is_file():
            yield p

def parse_top_blocks(text):
    out = {}
    i=0
    pat=re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{')
    n=len(text)
    while True:
        m=pat.search(text,i)
        if not m: break
        k=m.group(1); s=m.end()-1; d=0; j=s
        while j<n:
            c=text[j]
            if c=='{': d+=1
            elif c=='}':
                d-=1
                if d==0:
                    out[k]=text[s+1:j]
                    i=j+1
                    break
            j+=1
        else: break
    return out

def load_blocks(folder):
    d={}
    for f in iter_files(folder):
        d.update(parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')))
    return d

def parse_loc(folder):
    loc={}
    for f in folder.rglob('*.yml'):
        t=f.read_text(encoding='utf-8-sig', errors='ignore')
        for line in t.splitlines():
            s=line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'): continue
            m=re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$',s)
            if m: loc[m.group(1)] = m.group(2).replace('\\"','"').strip()
    return loc

def parse_comment_names(folder):
    out={}
    pat=re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{\s*#\s*(.+?)\s*$')
    for f in iter_files(folder):
        t=f.read_text(encoding='utf-8-sig', errors='ignore')
        for m in pat.finditer(t):
            k=m.group(1)
            v=m.group(2).strip()
            v=re.split(r'\btexture\b\s*=|\t', v, maxsplit=1)[0].strip()
            if v and k not in out: out[k]=v
    return out

def load_costs():
    c={}
    for k,b in parse_top_blocks(van_goods.read_text(encoding='utf-8-sig', errors='ignore')).items():
        m=re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)',b)
        if m: c[k]=float(m.group(1))
    for f in iter_files(mod_goods_dir):
        for k,b in parse_top_blocks(f.read_text(encoding='utf-8-sig', errors='ignore')).items():
            m=re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)',b)
            if m: c[k]=float(m.group(1))
    return c

def extract_io(body):
    ins, outs = {}, {}
    for m in re.finditer(r'(?m)^\s*(?!#)goods_input_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g,v=m.group(1),float(m.group(2)); ins[g]=ins.get(g,0.0)+v
    for m in re.finditer(r'(?m)^\s*(?!#)goods_output_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g,v=m.group(1),float(m.group(2)); outs[g]=outs.get(g,0.0)+v
    return ins, outs

loc=parse_loc(loc_dir)
comments={}
comments.update(parse_comment_names(mod_bld_dir)); comments.update(parse_comment_names(mod_pmg_dir)); comments.update(parse_comment_names(mod_pm_dir))

pm_v=load_blocks(van_pm_dir)
pm_m=load_blocks(mod_pm_dir)
pm_all=dict(pm_v); pm_all.update(pm_m)
new_pm=set(k for k in pm_m if k not in pm_v)

pmg_all=load_blocks(van_pmg_dir); pmg_all.update(load_blocks(mod_pmg_dir))
bld_all=load_blocks(van_bld_dir); bld_all.update(load_blocks(mod_bld_dir))

# ordered lists
pmg_order=[]
pmg_seen=set()
# collect by scanning files (van then mod so mod append)
for folder in [van_pmg_dir, mod_pmg_dir]:
    for f in iter_files(folder):
        txt=f.read_text(encoding='utf-8-sig', errors='ignore')
        for k in re.findall(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{', txt):
            if k in pmg_all and k not in pmg_seen:
                pmg_seen.add(k); pmg_order.append(k)

pmg_index={k:i for i,k in enumerate(pmg_order)}

pmg_to_pm_ordered={}
pm_to_pmgs=defaultdict(list)
for pmg, body in pmg_all.items():
    m=re.search(r'(?ms)production_methods\s*=\s*\{(.*?)\}', body)
    arr=[]
    if m:
        for ln in m.group(1).splitlines():
            s=ln.strip()
            if not s or s.startswith('#'): continue
            s=s.split('#',1)[0].strip()
            if re.match(r'^[A-Za-z0-9_]+$',s): arr.append(s)
    pmg_to_pm_ordered[pmg]=arr
    for idx,pm in enumerate(arr):
        pm_to_pmgs[pm].append((pmg, idx))

# PMG -> building(s) with order from building file
bld_order=[]; bseen=set()
for folder in [van_bld_dir, mod_bld_dir]:
    for f in iter_files(folder):
        txt=f.read_text(encoding='utf-8-sig', errors='ignore')
        for k in re.findall(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{', txt):
            if k in bld_all and k not in bseen:
                bseen.add(k); bld_order.append(k)

bld_index={k:i for i,k in enumerate(bld_order)}
pmg_to_blds=defaultdict(list)
for bld, body in bld_all.items():
    m=re.search(r'(?ms)production_method_groups\s*=\s*\{(.*?)\}', body)
    if not m: continue
    for pos,ln in enumerate(m.group(1).splitlines()):
        s=ln.strip()
        if not s or s.startswith('#'): continue
        s=s.split('#',1)[0].strip()
        if re.match(r'^[A-Za-z0-9_]+$',s):
            pmg_to_blds[s].append((bld,pos))

# manual bridge for commented railway pmg
manual_pmg='pmg_magic_railway'
for pm,idx in [('pm_no_magic',0),('pm_railway_luotuomu',1),('pm_railway_transport',2)]:
    if pm in new_pm:
        pm_to_pmgs[pm].append((manual_pmg,idx))
if manual_pmg not in pmg_index:
    pmg_index[manual_pmg]=10**8
if manual_pmg not in pmg_to_blds:
    pmg_to_blds[manual_pmg]=[('building_railway',0)]
comments.setdefault('pmg_magic_railway','魔法铁路方案')

costs=load_costs()
goods_alias={'magictools':loc.get('magic_tools','魔导器'),'magic_mineral':loc.get('magic_mineral','魔法矿石'),'magic_energy':loc.get('magic_energy','魔力'),'magic_steels':loc.get('magic_steels','超冶金结构'),'omocha':loc.get('omocha','周边')}

def cn(k): return loc.get(k) or comments.get(k) or {'building_railway':'铁路','pm_no_magic':'不启用','pm_railway_luotuomu':'洛托姆货运','pm_railway_transport':'反重力附魔'}.get(k,k)
def gcn(g): return loc.get(g) or goods_alias.get(g) or g

rows=[]
for pm in new_pm:
    ins,outs=extract_io(pm_all.get(pm,''))
    in_total=sum(q*costs.get(g,0.0) for g,q in ins.items())
    out_total=sum(q*costs.get(g,0.0) for g,q in outs.items())
    net=out_total-in_total
    roi=(net/in_total) if in_total else None
    maps=pm_to_pmgs.get(pm)
    if not maps: maps=[('(未匹配生产方式组)',10**8)]
    for pmg,pm_pos in maps:
        bmaps=pmg_to_blds.get(pmg)
        if not bmaps: bmaps=[('(未匹配建筑)',10**8)]
        for bld,bpos in bmaps:
            rows.append({
                'bld':bld,'pmg':pmg,'pm':pm,
                'bld_cn':cn(bld),'pmg_cn':cn(pmg),'pm_cn':cn(pm),
                'pm_pos':pm_pos,'bld_pmg_pos':bpos,
                'in':'；'.join([f"{gcn(g)}*{q:g}" for g,q in sorted(ins.items())]),
                'in_total':in_total,
                'out':'；'.join([f"{gcn(g)}*{q:g}" for g,q in sorted(outs.items())]),
                'out_total':out_total,'net':net,'roi':roi,
            })

rows.sort(key=lambda r:(bld_index.get(r['bld'],10**8), pmg_index.get(r['pmg'],10**8), r['bld_pmg_pos'], r['pm_pos'], r['pm_cn']))

out=root/'analysis'/'mod_new_pm_cn_grouped_v6_顺序修正.xlsx'
wb=Workbook(); ws=wb.active; ws.title='建筑-组-方式'
ws.append(['建筑','生产方式组','生产方式','商品投入（中文）','商品投入总价格','商品产出（中文）','商品产出总价格','商品收益净值','收益率'])
for r in rows:
    ws.append([r['bld_cn'],r['pmg_cn'],r['pm_cn'],r['in'],r['in_total'],r['out'],r['out_total'],r['net'],r['roi']])
for c in ws[1]: c.font = c.font.copy(bold=True)
for rr in ws.iter_rows(min_row=2,min_col=5,max_col=8):
    for c in rr: c.number_format='0.00'
for rr in ws.iter_rows(min_row=2,min_col=9,max_col=9):
    for c in rr: c.number_format='0.00%'
for i,w in {1:24,2:24,3:24,4:62,5:14,6:62,7:14,8:14,9:10}.items(): ws.column_dimensions[chr(64+i)].width=w
wb.save(out)
print(out)
