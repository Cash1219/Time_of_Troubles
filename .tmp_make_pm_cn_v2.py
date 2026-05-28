import re
from pathlib import Path
from collections import defaultdict
from openpyxl import Workbook

root = Path(r"D:/Users/Administrator/Documents/Paradox Interactive/Victoria 3/mod/ToT")
van_root = Path(r"D:/Steam/steamapps/common/Victoria 3/game")

loc_dir = root/'localization'/'simp_chinese'
mod_pm_dir = root/'common'/'production_methods'
van_pm_dir = van_root/'common'/'production_methods'
mod_pmg_dir = root/'common'/'production_method_groups'
van_pmg_dir = van_root/'common'/'production_method_groups'
mod_bld_dir = root/'common'/'buildings'
van_bld_dir = van_root/'common'/'buildings'
van_goods = van_root/'common'/'goods'/'00_goods.txt'
mod_goods_dir = root/'common'/'goods'


def iter_text_files(folder: Path):
    if not folder.exists():
        return
    for p in folder.rglob('*'):
        if p.is_file():
            yield p


def parse_top_blocks(text):
    blocks = {}
    i = 0
    pat = re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{')
    n = len(text)
    while True:
        m = pat.search(text, i)
        if not m:
            break
        key = m.group(1)
        s = m.end() - 1
        d = 0
        j = s
        while j < n:
            c = text[j]
            if c == '{': d += 1
            elif c == '}':
                d -= 1
                if d == 0:
                    blocks[key] = text[s+1:j]
                    i = j + 1
                    break
            j += 1
        else:
            break
    return blocks


def parse_loc(folder):
    loc = {}
    for f in folder.rglob('*.yml'):
        txt = f.read_text(encoding='utf-8-sig', errors='ignore')
        for line in txt.splitlines():
            s = line.strip()
            if not s or s.startswith('#') or s.startswith('l_simp_chinese:'):
                continue
            m = re.match(r'^([A-Za-z0-9_\.\-]+)\s*:\s*"(.*)"\s*$', s)
            if m:
                loc[m.group(1)] = m.group(2).replace('\\"','"').strip()
    return loc


def parse_key_comment_name_map(folder):
    # capture inline Chinese comments like: key = { #中文名
    out = {}
    pat = re.compile(r'(?m)^\s*([A-Za-z0-9_]+)\s*=\s*\{\s*#\s*(.+?)\s*$')
    for f in iter_text_files(folder):
        txt = f.read_text(encoding='utf-8-sig', errors='ignore')
        for m in pat.finditer(txt):
            key = m.group(1)
            name = m.group(2).strip()
            if name and key not in out:
                out[key] = name
    return out


def load_blocks(folder):
    out = {}
    for f in iter_text_files(folder):
        txt = f.read_text(encoding='utf-8-sig', errors='ignore')
        out.update(parse_top_blocks(txt))
    return out


def load_goods_costs():
    costs = {}
    for k,body in parse_top_blocks(van_goods.read_text(encoding='utf-8-sig', errors='ignore')).items():
        m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', body)
        if m: costs[k] = float(m.group(1))
    if mod_goods_dir.exists():
        for f in iter_text_files(mod_goods_dir):
            txt = f.read_text(encoding='utf-8-sig', errors='ignore')
            for k,body in parse_top_blocks(txt).items():
                m = re.search(r'(?m)^\s*cost\s*=\s*([0-9]+(?:\.[0-9]+)?)', body)
                if m: costs[k] = float(m.group(1))
    return costs


def extract_io(body):
    ins, outs = {}, {}
    for m in re.finditer(r'(?m)^\s*(?!#)goods_input_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g, v = m.group(1), float(m.group(2)); ins[g] = ins.get(g,0.0)+v
    for m in re.finditer(r'(?m)^\s*(?!#)goods_output_([A-Za-z0-9_]+)_add\s*=\s*([+-]?[0-9]+(?:\.[0-9]+)?)', body):
        g, v = m.group(1), float(m.group(2)); outs[g] = outs.get(g,0.0)+v
    return ins, outs

# load base data
loc = parse_loc(loc_dir)
comment_names = {}
comment_names.update(parse_key_comment_name_map(mod_pm_dir))
comment_names.update(parse_key_comment_name_map(mod_pmg_dir))
comment_names.update(parse_key_comment_name_map(mod_bld_dir))

pm_blocks_van = load_blocks(van_pm_dir)
pm_blocks_mod = load_blocks(mod_pm_dir)
pm_blocks = dict(pm_blocks_van)
pm_blocks.update(pm_blocks_mod)
new_pm_keys = sorted([k for k in pm_blocks_mod.keys() if k not in pm_blocks_van])

pmg_blocks = load_blocks(van_pmg_dir)
pmg_blocks.update(load_blocks(mod_pmg_dir))

bld_blocks = load_blocks(van_bld_dir)
bld_blocks.update(load_blocks(mod_bld_dir))

# map PM -> PMG via production_methods list
pm_to_pmg = defaultdict(set)
for pmg_key, body in pmg_blocks.items():
    m = re.search(r'(?ms)production_methods\s*=\s*\{(.*?)\}', body)
    if not m:
        continue
    section = m.group(1)
    for line in section.splitlines():
        s = line.strip()
        if not s or s.startswith('#'):
            continue
        s = s.split('#',1)[0].strip()
        if re.match(r'^[A-Za-z0-9_]+$', s):
            pm_to_pmg[s].add(pmg_key)

# map PMG -> Building via production_method_groups list
pmg_to_bld = defaultdict(set)
for bld_key, body in bld_blocks.items():
    m = re.search(r'(?ms)production_method_groups\s*=\s*\{(.*?)\}', body)
    if not m:
        continue
    section = m.group(1)
    for line in section.splitlines():
        s = line.strip()
        if not s or s.startswith('#'):
            continue
        s = s.split('#',1)[0].strip()
        if re.match(r'^[A-Za-z0-9_]+$', s):
            pmg_to_bld[s].add(bld_key)

costs = load_goods_costs()

# goods chinese alias fallback
goods_alias = {
    'magictools': loc.get('magic_tools','魔导器'),
    'magic_mineral': loc.get('magic_mineral','魔法矿石'),
    'magic_energy': loc.get('magic_energy','魔力'),
    'magic_steels': loc.get('magic_steels','超冶金结构'),
    'omocha': loc.get('omocha','周边'),
}

def cn_name(key):
    return loc.get(key) or comment_names.get(key) or key

def goods_cn(g):
    return loc.get(g) or goods_alias.get(g) or g

rows = []
for pm in new_pm_keys:
    ins, outs = extract_io(pm_blocks.get(pm,''))
    in_total = sum(q*costs.get(g,0.0) for g,q in ins.items())
    out_total = sum(q*costs.get(g,0.0) for g,q in outs.items())
    net = out_total - in_total
    roi = (net/in_total) if in_total else None

    pmgs = sorted(pm_to_pmg.get(pm) or ['(未匹配生产方式组)'])
    for pmg in pmgs:
        blds = sorted(pmg_to_bld.get(pmg) or ['(未匹配建筑)'])
        for bld in blds:
            rows.append([
                cn_name(bld),
                cn_name(pmg),
                cn_name(pm),
                '；'.join([f"{goods_cn(g)}*{q:g}" for g,q in sorted(ins.items())]),
                in_total,
                '；'.join([f"{goods_cn(g)}*{q:g}" for g,q in sorted(outs.items())]),
                out_total,
                net,
                roi,
            ])

rows.sort(key=lambda r: (r[0], r[1], r[2]))

out = root/'analysis'/'mod_new_pm_cn_grouped_v2.xlsx'
wb = Workbook()
ws = wb.active
ws.title = '建筑-生产方式组-生产方式'
ws.append(['建筑','生产方式组','生产方式','商品投入（中文）','商品投入总价格','商品产出（中文）','商品产出总价格','商品收益净值','收益率'])
for r in rows:
    ws.append(r)

for c in ws[1]:
    c.font = c.font.copy(bold=True)
for row in ws.iter_rows(min_row=2, min_col=5, max_col=8):
    for c in row: c.number_format = '0.00'
for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
    for c in row: c.number_format = '0.00%'
for i,w in {1:24,2:24,3:24,4:62,5:14,6:62,7:14,8:14,9:10}.items():
    ws.column_dimensions[chr(64+i)].width = w
wb.save(out)

# quick validation
import re
unmatched = 0
english = 0
for r in rows:
    txt = ' | '.join([str(x or '') for x in [r[0],r[1],r[2],r[3],r[5]]])
    if '未匹配' in txt:
        unmatched += 1
    if re.search(r'[A-Za-z_]{3,}', txt):
        english += 1
print(out)
print(f"rows={len(rows)} unmatched={unmatched} english_like={english}")
